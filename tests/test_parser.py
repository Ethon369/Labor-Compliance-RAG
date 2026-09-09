"""文档解析的离线单测：构造字节流即可，不依赖真文件、真库、网络。

PDF 样本用代码手搓（含 xref 表），不往仓库塞二进制 fixture：样本是"测试数据"，
写清楚怎么生成的比塞一个看不懂的 .pdf 更可维护，也避免仓库里出现来源不明的文件。
"""
from __future__ import annotations

import io

import pytest

from app.kb.parser import (
    Block,
    ParsedDoc,
    UnsupportedContentError,
    detect_source_type,
    parse_file,
)


def _build_pdf(lines: list[str]) -> bytes:
    """生成一个最小的、带文字层的单页 PDF（Helvetica 内置字体，仅支持 latin-1）。"""
    stream = ("BT /F1 18 Tf 72 700 Td 20 TL\n"
              + "".join(f"({t}) Tj T*\n" for t in lines)
              + "ET\n").encode("latin-1")
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"endstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj".encode() + body + b"endobj\n"
    xref_off = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer<</Size {len(objs) + 1}/Root 1 0 R>>\nstartxref\n{xref_off}\n%%EOF\n".encode()
    return bytes(out)


def _write(tmp_path, name: str, data: bytes) -> str:
    p = tmp_path / name
    p.write_bytes(data)
    return str(p)


# ---- 类型识别 ----

@pytest.mark.parametrize("filename,expected", [
    ("a.pdf", "pdf"), ("a.PDF", "pdf"),
    ("b.docx", "docx"),
    ("c.txt", "txt"), ("d.text", "text"),
    ("e.md", "md"), ("f.markdown", "md"),
])
def test_detect_source_type(filename, expected):
    assert detect_source_type(filename) == expected


def test_detect_source_type_rejects_unknown():
    assert detect_source_type("virus.exe") is None
    assert detect_source_type("noext") is None


# ---- 纯文本 / Markdown ----

def test_parse_text_utf8(tmp_path):
    path = _write(tmp_path, "a.txt", "第一条 内容甲\n\n第二条 内容乙\n".encode("utf-8"))
    doc = parse_file(path, "txt")
    assert [b.text for b in doc.blocks] == ["第一条 内容甲", "第二条 内容乙"]
    assert all(not b.is_heading for b in doc.blocks)
    assert doc.page_count is None


def test_parse_text_gb18030_fallback(tmp_path):
    """老中文文档常见 GBK/GB18030 编码：直接按 utf-8 解码会抛，必须能回退。"""
    path = _write(tmp_path, "gbk.txt", "第一条 中文内容".encode("gb18030"))
    doc = parse_file(path, "txt")
    assert doc.blocks[0].text == "第一条 中文内容"


def test_parse_markdown_marks_headings(tmp_path):
    path = _write(tmp_path, "a.md", "# 第一章 总则\n第一条 内容\n## 小节\n".encode("utf-8"))
    doc = parse_file(path, "md")
    assert [(b.text, b.is_heading) for b in doc.blocks] == [
        ("# 第一章 总则", True),
        ("第一条 内容", False),
        ("## 小节", True),
    ]


def test_parse_empty_text_raises(tmp_path):
    path = _write(tmp_path, "empty.txt", b"   \n\n  ")
    with pytest.raises(UnsupportedContentError):
        parse_file(path, "txt")


def test_parse_unknown_type_raises(tmp_path):
    path = _write(tmp_path, "x.bin", b"data")
    with pytest.raises(UnsupportedContentError):
        parse_file(path, "exe")


# ---- Word ----

def test_parse_docx_marks_heading_styles(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("第一章 总则", level=1)
    d.add_paragraph("第一条 为了公正及时解决劳动争议。")
    path = str(tmp_path / "a.docx")
    d.save(path)

    doc = parse_file(path, "docx")
    assert [(b.text, b.is_heading) for b in doc.blocks] == [
        ("第一章 总则", True),
        ("第一条 为了公正及时解决劳动争议。", False),
    ]


def test_parse_empty_docx_raises(tmp_path):
    import docx

    path = str(tmp_path / "empty.docx")
    docx.Document().save(path)
    with pytest.raises(UnsupportedContentError):
        parse_file(path, "docx")


# ---- PDF ----

def test_parse_pdf_keeps_page_number(tmp_path):
    path = _write(tmp_path, "a.pdf", _build_pdf(["Article 1 Enacted", "Article 2 Scope"]))
    doc = parse_file(path, "pdf")
    assert doc.page_count == 1
    assert [b.text for b in doc.blocks] == ["Article 1 Enacted", "Article 2 Scope"]
    assert all(b.page == 1 for b in doc.blocks)


def test_parse_pdf_without_text_layer_raises(tmp_path):
    """扫描版 PDF 没有文字层：解析器要抛可识别的异常，让上层标 failed 而不是炸整条流水线。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    path = _write(tmp_path, "scan.pdf", buf.getvalue())

    with pytest.raises(UnsupportedContentError) as e:
        parse_file(path, "pdf")
    assert "OCR" in str(e.value)


# ---- 数据模型 ----

def test_parsed_doc_char_count():
    doc = ParsedDoc(blocks=[Block(text="abc"), Block(text="de", page=1)])
    assert doc.char_count == 5
