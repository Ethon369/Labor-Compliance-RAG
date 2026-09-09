"""入库流水线集成测试（@pytest.mark.db）。

核心契约只有一条：**run() 永不抛异常**——它是后台任务，抛出去了没人接，文档会永远
卡在 pending。失败必须落成 status='failed' + error 有内容。
"""
from __future__ import annotations

import secrets

import psycopg
import pytest

from app.api.auth import ensure_user_schema, hash_password
from app.core.config import Settings
from app.kb import pipeline, store
from app.kb.parser import UnsupportedContentError
from app.kb.schema import ensure_kb_schema
from app.rag.vectorstore import PgVectorStore

pytestmark = pytest.mark.db


@pytest.fixture
def env(tmp_path):
    s = Settings()
    try:
        conn = psycopg.connect(s.database_url)
    except psycopg.OperationalError as e:
        pytest.skip(f"labor-pg 不可达（docker compose up -d 后重跑）：{e}")
    ensure_user_schema(s.database_url)
    ensure_kb_schema(s.database_url)
    PgVectorStore(s.database_url, s.embedding_dim).ensure_ready()

    uid = conn.execute(
        "INSERT INTO users (username, password_hash) VALUES (%s, %s) RETURNING id",
        ("u_" + secrets.token_hex(4), hash_password("x")),
    ).fetchone()[0]
    conn.commit()
    conn.close()
    return s.database_url, uid, tmp_path


def _new_kb(dsn, uid, **kw) -> dict:
    kb_id = store.create_kb(dsn, uid, "kb_" + secrets.token_hex(3), **kw)
    return store.fetch_kb(dsn, kb_id)


def _new_doc(dsn, kb_id, path, source_type) -> int:
    return store.create_document(dsn, kb_id, path.name, source_type,
                                 file_size=path.stat().st_size)


def _write(tmp_path, name: str, text: str):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


# ---- 成功路径 ----

def test_txt_document_becomes_ready_with_vectors(env):
    dsn, uid, tmp = env
    kb = _new_kb(dsn, uid)
    path = _write(tmp, "手册.txt", "第一章 总则\n" + "甲" * 400 + "。\n" + "乙" * 400 + "。")
    doc_id = _new_doc(dsn, kb["id"], path, "txt")

    pipeline.run(dsn, doc_id, path, "txt", kb)

    doc = store.get_document(dsn, doc_id)
    assert doc["status"] == "ready" and doc["error"] is None
    assert doc["chunk_count"] > 1
    assert store.fetch_kb(dsn, kb["id"])["chunk_count"] == doc["chunk_count"]
    # 内容与向量同事务写入，ready 时不该存在"没有向量的片"
    with psycopg.connect(dsn) as c:
        missing = c.execute(
            "SELECT count(*) FROM chunks WHERE doc_id = %s AND embedding IS NULL", (doc_id,)
        ).fetchone()[0]
    assert missing == 0
    assert not path.exists()          # 临时文件已清理


def test_docx_document_is_parsed_end_to_end(env):
    dsn, uid, tmp = env
    import docx

    kb = _new_kb(dsn, uid)
    path = tmp / "制度.docx"
    document = docx.Document()
    document.add_paragraph("员工手册", style="Heading 1")
    document.add_paragraph("甲" * 300 + "。")
    document.save(str(path))

    doc_id = _new_doc(dsn, kb["id"], path, "docx")
    pipeline.run(dsn, doc_id, path, "docx", kb)

    doc = store.get_document(dsn, doc_id)
    assert doc["status"] == "ready" and doc["chunk_count"] >= 1
    with psycopg.connect(dsn) as c:
        row = c.execute("SELECT heading FROM chunks WHERE doc_id = %s ORDER BY seq",
                        (doc_id,)).fetchone()
    assert row[0] == "员工手册"      # DOCX 的 Heading 样式成为切片的上下文标题


def test_chunk_params_come_from_kb_row(env):
    """切分粒度是库的属性：同一份文档进不同参数的库，片数应该不同。"""
    dsn, uid, tmp = env
    text = "甲" * 900 + "。"
    small = _new_kb(dsn, uid, chunk_size=100, chunk_overlap=0)
    big = _new_kb(dsn, uid, chunk_size=500, chunk_overlap=0)

    counts = {}
    for name, kb in (("small", small), ("big", big)):
        path = _write(tmp, f"{name}.txt", text)
        doc_id = _new_doc(dsn, kb["id"], path, "txt")
        pipeline.run(dsn, doc_id, path, "txt", kb)
        counts[name] = store.get_document(dsn, doc_id)["chunk_count"]

    assert counts["small"] > counts["big"] > 0


# ---- 失败路径（run 不抛，落 failed）----

def test_unsupported_content_becomes_failed_not_exception(env, monkeypatch):
    """扫描件 PDF：文件能开但没文字层 → failed + 明确原因，而不是后台炸掉。"""
    dsn, uid, tmp = env
    kb = _new_kb(dsn, uid)
    path = _write(tmp, "扫描件.pdf", "%PDF-1.4 假文件")
    doc_id = _new_doc(dsn, kb["id"], path, "pdf")

    def _boom(_p, _t):
        raise UnsupportedContentError("PDF 无文字层（可能是扫描件），需要 OCR，当前版本不支持")

    monkeypatch.setattr(pipeline, "parse_file", _boom)
    pipeline.run(dsn, doc_id, path, "pdf", kb)      # 不得抛出

    doc = store.get_document(dsn, doc_id)
    assert doc["status"] == "failed"
    assert "OCR" in (doc["error"] or "")


def test_unexpected_exception_is_captured_into_error(env, monkeypatch):
    dsn, uid, tmp = env
    kb = _new_kb(dsn, uid)
    path = _write(tmp, "坏文件.txt", "内容")
    doc_id = _new_doc(dsn, kb["id"], path, "txt")

    def _boom(_p, _t):
        raise RuntimeError("磁盘读到一半坏了")

    monkeypatch.setattr(pipeline, "parse_file", _boom)
    pipeline.run(dsn, doc_id, path, "txt", kb)      # 不得抛出

    doc = store.get_document(dsn, doc_id)
    assert doc["status"] == "failed"
    assert "RuntimeError" in (doc["error"] or "") and "坏了" in doc["error"]


def test_blank_file_is_failed_with_readable_reason(env):
    """空白文件在解析阶段就被拦下：failed + 人能读懂的原因，而不是 chunk_count=0 的 ready。"""
    dsn, uid, tmp = env
    kb = _new_kb(dsn, uid)
    path = _write(tmp, "空白.txt", "   \n  \n")
    doc_id = _new_doc(dsn, kb["id"], path, "txt")

    pipeline.run(dsn, doc_id, path, "txt", kb)

    doc = store.get_document(dsn, doc_id)
    assert doc["status"] == "failed" and doc["chunk_count"] == 0
    assert "空" in (doc["error"] or "")


def test_zero_chunks_is_failed_instead_of_empty_ready(env, monkeypatch):
    """防御分支：解析出了块却切不出片（异常输入）——宁可 failed，不留空 ready 文档。"""
    dsn, uid, tmp = env
    kb = _new_kb(dsn, uid)
    path = _write(tmp, "x.txt", "内容")
    doc_id = _new_doc(dsn, kb["id"], path, "txt")

    monkeypatch.setattr(pipeline, "split_document", lambda *a, **kw: [])
    pipeline.run(dsn, doc_id, path, "txt", kb)

    doc = store.get_document(dsn, doc_id)
    assert doc["status"] == "failed"
    assert "没有可入库" in (doc["error"] or "")


def test_missing_document_row_is_ignored(env):
    """文档行被并发删除：流水线不该把异常抛进后台任务，也不该建出孤儿切片。"""
    dsn, uid, tmp = env
    kb = _new_kb(dsn, uid)
    path = _write(tmp, "x.txt", "内容")
    pipeline.run(dsn, 987654, path, "txt", kb)      # 不存在的 doc_id
    assert not path.exists()


def test_reprocess_reuses_unchanged_embeddings(env, monkeypatch):
    """同一文档重处理时，内容未变的切片复用旧向量，不再调用 embedding。

    嵌入去重（V2.4）：pipeline 写前查 (seq, content_hash)，匹配就复用旧向量，
    只对新增/变更切片重新向量化。这里用计数 embedder 验证第二次 embed 调用为 0。
    """
    from app.rag.vectorstore import parse_vector

    dsn, uid, tmp = env
    kb = _new_kb(dsn, uid)
    text = "第一章 总则\n" + "甲" * 400 + "。\n" + "乙" * 400 + "。"

    real_build = pipeline.build_embedder
    calls = {"n": 0}

    def counting_build(settings):
        emb = real_build(settings)

        class _Counting:
            def embed(self, texts):
                calls["n"] += len(texts)
                return emb.embed(texts)
        return _Counting()

    monkeypatch.setattr(pipeline, "build_embedder", counting_build)

    # 第一次：全新内容，全量 embed
    path1 = _write(tmp, "a.txt", text)
    doc_id = _new_doc(dsn, kb["id"], path1, "txt")
    pipeline.run(dsn, doc_id, path1, "txt", kb)
    first_calls = calls["n"]
    assert first_calls > 0

    # 第二次：同 doc_id 重处理，内容逐字相同 → 全部复用，embed 调用不增加
    path2 = _write(tmp, "b.txt", text)
    pipeline.run(dsn, doc_id, path2, "txt", kb)
    assert calls["n"] == first_calls

    # 复用后向量仍在且非空（不是 NULL 的"待回填"态）
    embedded = store.doc_embedded(dsn, doc_id)
    assert len(embedded) == first_calls
    assert all(parse_vector(v) for _, (_, v) in embedded.items())


def test_reprocess_only_reembeds_changed_slices(env, monkeypatch):
    """重处理时只有内容变化的切片重新向量化，未变的仍复用。

    用 chunk_overlap=0 让切片内容完全独立：改第一段不会通过 overlap 尾缀污染第二片，
    这样"只重算变化切片"的断言才成立。
    """
    dsn, uid, tmp = env
    kb = _new_kb(dsn, uid, chunk_overlap=0)
    text = "第一章 总则\n" + "甲" * 400 + "。\n" + "乙" * 400 + "。"

    real_build = pipeline.build_embedder
    calls = {"n": 0}

    def counting_build(settings):
        emb = real_build(settings)

        class _Counting:
            def embed(self, texts):
                calls["n"] += len(texts)
                return emb.embed(texts)
        return _Counting()

    monkeypatch.setattr(pipeline, "build_embedder", counting_build)

    path1 = _write(tmp, "a.txt", text)
    doc_id = _new_doc(dsn, kb["id"], path1, "txt")
    pipeline.run(dsn, doc_id, path1, "txt", kb)
    first_calls = calls["n"]

    # 只改第一段内容（甲→丙），第二段不变
    changed = "第一章 总则\n" + "丙" * 400 + "。\n" + "乙" * 400 + "。"
    path2 = _write(tmp, "b.txt", changed)
    pipeline.run(dsn, doc_id, path2, "txt", kb)
    # 变化切片重新 embed，未变切片复用：新增 embed 调用应 < 总片数
    assert 0 < (calls["n"] - first_calls) < first_calls
